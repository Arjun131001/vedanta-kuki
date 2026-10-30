import io
from dataclasses import replace

import openpyxl
from fastapi.testclient import TestClient

from app.config import Settings
from app.main import create_app
from app.store import Store
from app.worker import process_one


def empty_source():
    return {"version": 1, "kind": "xlsx", "summary": {"warnings": []}, "units": []}


def settings(tmp_path, **kwargs):
    return replace(Settings(), data_dir=tmp_path, **kwargs)


def workbook():
    book = openpyxl.Workbook()
    book.active["A1"] = "Scope 1"
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def upload(client, payload=None, name="report.xlsx"):
    return client.post("/api/documents", files=[("files", (name, payload if payload is not None else workbook()))])


def test_persist_deduplicate_and_validate_after_restart(tmp_path):
    config = settings(tmp_path)
    payload = workbook()
    with TestClient(create_app(config)) as client:
        created = upload(client, payload).json()["results"][0]
        document_id = created["document"]["id"]
        assert created["document"]["status"] == "queued"
        duplicate = upload(client, payload, "different-name.xlsx").json()["results"][0]
        assert duplicate["duplicate"] is True
        assert duplicate["document"]["id"] == document_id
    with TestClient(create_app(config)) as client:
        assert len(client.get("/api/documents").json()["documents"]) == 1
        assert process_one(Store(config))
        result = client.get(f"/api/documents/{document_id}").json()
        assert result["status"] == "ready"
        assert result["validation"]["sheet_count"] == 1
        assert result["validation"]["metrics_extracted"] is False
        assert next(config.uploads_dir.iterdir()).read_bytes() == payload


def test_retry_backoff_and_terminal_failure(tmp_path):
    config = settings(tmp_path, retry_base_seconds=5)
    with TestClient(create_app(config)) as client:
        doc = upload(client, b"broken", "bad.pdf").json()["results"][0]["document"]
        store = Store(config)
        first = store.claim(now=doc["created_at"] + 1)
        store.finish(first, error="unreadable", now=100)
        assert store.claim(now=104) is None
        second = store.claim(now=105)
        assert second["attempts"] == 2
        store.finish(second, error="unreadable", now=105)
        assert store.claim(now=114) is None
        third = store.claim(now=115)
        store.finish(third, error="unreadable", now=115)
        assert store.get(doc["id"])["status"] == "failed"
        assert client.post(f'/api/documents/{doc["id"]}/retry').status_code == 200
        assert store.get(doc["id"])["attempts"] == 0
        assert client.post(f'/api/documents/{doc["id"]}/retry').status_code == 409


def test_lease_recovery_and_stale_completion(tmp_path):
    config = settings(tmp_path, lease_seconds=10)
    with TestClient(create_app(config)) as client:
        doc = upload(client).json()["results"][0]["document"]
    store = Store(config)
    now = doc["created_at"] + 1
    abandoned = store.claim(now=now)
    assert Store(config).claim(now=now + 9) is None
    recovered = Store(config).claim(now=now + 11)
    assert recovered["attempts"] == 2
    assert store.finish(abandoned, validation={"wrong": True}, source=empty_source()) is False
    assert store.finish(recovered, validation={"sheet_count": 1}, source=empty_source()) is True
    assert store.get(doc["id"])["status"] == "ready"


def test_real_parser_failure_retries_are_bounded(tmp_path):
    config = settings(tmp_path, retry_base_seconds=0)
    with TestClient(create_app(config)) as client:
        doc = upload(client, b"not a workbook").json()["results"][0]["document"]
        store = Store(config)
        assert process_one(store)
        result = store.get(doc["id"])
        assert result["status"] == "failed"
        assert result["attempts"] == 1
        assert "valid XLSX" in result["error"]
        assert process_one(store) is False


def test_upload_limits_safe_names_and_partial_batch(tmp_path):
    config = settings(tmp_path, max_upload_bytes=12)
    with TestClient(create_app(config)) as client:
        results = client.post("/api/documents", files=[
            ("files", ("../../bad.pdf", b"%PDF-data")),
            ("files", ("huge.pdf", b"x" * 13)),
            ("files", ("notes.txt", b"hello")),
            ("files", ("empty.pdf", b"")),
        ]).json()["results"]
        assert results[0]["document"]["filename"] == "bad.pdf"
        assert "storage_key" not in results[0]["document"]
        assert all("error" in row for row in results[1:])
        assert not list(tmp_path.glob("incoming-*"))
        assert client.get("/api/documents/missing").status_code == 404
        assert client.get("/api/health").json()["status"] == "ok"


def test_simultaneous_workers_claim_only_once(tmp_path):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier

    config = settings(tmp_path)
    with TestClient(create_app(config)) as client:
        upload(client)
    barrier = Barrier(2)
    def claim():
        store = Store(config)
        barrier.wait()
        return store.claim()
    with ThreadPoolExecutor(max_workers=2) as pool:
        claimed = list(pool.map(lambda _: claim(), range(2)))
    assert sum(job is not None for job in claimed) == 1


def test_expired_token_cannot_finish_even_before_reclaim(tmp_path):
    config = settings(tmp_path, lease_seconds=10)
    with TestClient(create_app(config)) as client:
        doc = upload(client).json()["results"][0]["document"]
    store = Store(config)
    now = doc["created_at"] + 1
    job = store.claim(now=now)
    assert not store.finish(job, validation={"wrong": True}, now=now + 11, source=empty_source())
    assert store.get(doc["id"])["status"] == "processing"


def pdf_document():
    # Minimal one-page PDF with a correctly calculated cross-reference table.
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] >>",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, content in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{number} 0 obj\n".encode() + content + b"\nendobj\n")
    xref = len(data)
    data.extend(b"xref\n0 4\n0000000000 65535 f \n")
    for offset in offsets[1:]:
        data.extend(f"{offset:010} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size 4 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(data)


def test_valid_pdf_and_transient_recovery(tmp_path):
    config = settings(tmp_path, retry_base_seconds=0)
    with TestClient(create_app(config)) as client:
        doc = upload(client, pdf_document(), "report.pdf").json()["results"][0]["document"]
        store = Store(config)
        def unavailable(path, kind):
            raise OSError("temporary read error")
        assert process_one(store, unavailable)
        assert store.get(doc["id"])["status"] == "queued"
        assert process_one(store)
        result = store.get(doc["id"])
        assert result["status"] == "ready"
        assert result["attempts"] == 2
        assert result["validation"]["page_count"] == 1


def test_source_workbook_formulas_pagination_and_persistence(tmp_path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Emissions"
    sheet["A1"] = "Scope 1"
    sheet["B1"] = 42
    sheet["C1"] = "=B1*2"
    sheet["B1"].number_format = '0.00'
    buffer = io.BytesIO()
    book.save(buffer)
    config = settings(tmp_path)
    with TestClient(create_app(config)) as client:
        doc = upload(client, buffer.getvalue()).json()["results"][0]["document"]
        assert client.get(f'/api/documents/{doc["id"]}/sources').status_code == 409
        assert process_one(Store(config))
    with TestClient(create_app(config)) as client:
        result = client.get(f'/api/documents/{doc["id"]}/sources?unit=1&offset=1&limit=2').json()
        assert result["units"] == [{"index": 1, "label": "Emissions"}]
        assert result["unit"]["total_cells"] == 3
        value, formula = result["unit"]["cells"]
        assert (value["coordinate"], value["value"], value["number_format"]) == ("B1", 42, "0.00")
        assert formula["coordinate"] == "C1" and formula["formula"] == "=B1*2"
        assert formula["cached_value"] is None
        assert "not recalculated" in result["summary"]["warnings"][0]
        assert client.get(f'/api/documents/{doc["id"]}/sources?unit=2').status_code == 404


def test_blank_pdf_page_preserved_and_warned(tmp_path):
    config = settings(tmp_path)
    with TestClient(create_app(config)) as client:
        doc = upload(client, pdf_document(), "blank.pdf").json()["results"][0]["document"]
        assert process_one(Store(config))
        result = client.get(f'/api/documents/{doc["id"]}/sources?unit=1').json()
        assert result["unit"]["text"] == ""
        assert "OCR may be needed" in result["summary"]["warnings"][0]


def test_legacy_ready_document_is_requeued_once(tmp_path):
    config = settings(tmp_path)
    with TestClient(create_app(config)) as client:
        doc = upload(client).json()["results"][0]["document"]
    store = Store(config)
    with store.connect() as db:
        db.execute("UPDATE documents SET status='ready',validation_json=? WHERE id=?", ('{"sheet_count":1}', doc["id"]))
        db.execute("UPDATE jobs SET status='completed',attempts=3 WHERE document_id=?", (doc["id"],))
    migrated = Store(config)
    assert migrated.get(doc["id"])["status"] == "queued"
    assert migrated.get(doc["id"])["attempts"] == 0
    assert migrated.get(doc["id"])["validation"]["sheet_count"] == 1
    assert process_one(migrated)
    assert Store(config).get(doc["id"])["status"] == "ready"
    assert Store(config).sources(doc["id"])["version"] == 1


def test_stale_worker_cannot_persist_source(tmp_path):
    config = settings(tmp_path, lease_seconds=10)
    with TestClient(create_app(config)) as client:
        doc = upload(client).json()["results"][0]["document"]
    store = Store(config)
    now = doc["created_at"] + 1
    stale = store.claim(now=now)
    store.claim(now=now + 11)
    assert not store.finish(stale, validation={}, source=empty_source(), now=now + 12)
    assert store.sources(doc["id"]) is None


def test_source_limits_fail_without_partial_result(tmp_path, monkeypatch):
    import app.extraction as extraction
    monkeypatch.setattr(extraction, "MAX_CELLS", 0)
    config = settings(tmp_path)
    with TestClient(create_app(config)) as client:
        doc = upload(client).json()["results"][0]["document"]
        assert process_one(Store(config))
        assert client.get(f'/api/documents/{doc["id"]}').json()["status"] == "failed"
        assert client.get(f'/api/documents/{doc["id"]}/sources').status_code == 409


def text_pdf():
    stream = b"BT /F1 12 Tf 20 100 Td (Scope 1 emissions) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    for offset in offsets:
        data.extend(f"{offset:010} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(data)


def test_pdf_text_is_persisted_with_page_location(tmp_path):
    config = settings(tmp_path)
    with TestClient(create_app(config)) as client:
        doc = upload(client, text_pdf(), "text.pdf").json()["results"][0]["document"]
        assert process_one(Store(config))
        result = client.get(f'/api/documents/{doc["id"]}/sources?unit=1').json()
        assert result["document_id"] == doc["id"]
        assert result["unit"]["label"] == "Page 1"
        assert "Scope 1 emissions" in result["unit"]["text"]
        assert result["summary"]["warnings"] == []


def test_inaccurate_workbook_dimensions_do_not_silently_drop_cells(tmp_path):
    from zipfile import ZipFile, ZIP_DEFLATED
    book = openpyxl.Workbook()
    book.active["A1"] = "visible"
    book.active["C3"] = "hidden by bad dimension"
    original = io.BytesIO()
    book.save(original)
    altered = io.BytesIO()
    with ZipFile(original) as source, ZipFile(altered, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'ref="A1:C3"', b'ref="A1:A1"', 1)
            target.writestr(entry.filename, data)
    config = settings(tmp_path)
    with TestClient(create_app(config)) as client:
        doc = upload(client, altered.getvalue()).json()["results"][0]["document"]
        assert process_one(Store(config))
        result = client.get(f'/api/documents/{doc["id"]}').json()
        assert result["status"] == "failed"
        assert "dimensions omit cells" in result["error"]
        assert client.get(f'/api/documents/{doc["id"]}/sources').status_code == 409


def test_legacy_pdf_layout_requeues_once_without_deleting_old_source(tmp_path):
    config=settings(tmp_path)
    with TestClient(create_app(config)) as client:
        doc=upload(client,text_pdf(),"text.pdf").json()["results"][0]["document"]
        store=Store(config); assert process_one(store)
        with store.connect() as db:
            db.execute("UPDATE source_results SET summary_json=? WHERE document_id=?", ('{"warnings":[],"unit_count":1}',doc["id"]))
        migrated=Store(config)
        assert migrated.get(doc["id"])["status"] == "queued"
        assert migrated.sources(doc["id"],unit=1)["unit"]["text"]
        assert process_one(migrated)
        assert Store(config).get(doc["id"])["status"] == "ready"
        assert Store(config).sources(doc["id"])["summary"]["layout_version"] == "pdf-layout-v2"
