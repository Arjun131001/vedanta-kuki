import json
from dataclasses import replace
import httpx
import pytest
from fastapi.testclient import TestClient
from app.config import Settings
from app.facts import FactError, call_openai, FactResult, project_result
from app.main import create_app
from app.store import Store
from app import fact_queue
from app.worker import process_one, process_fact
from test_intake import upload

SOURCE = {"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":"Acme Ltd emissions 1,234 tonnes in 2024."}]}
EVIDENCE = {"unit_index":1,"quote":"Acme Ltd emissions 1,234 tonnes in 2024.","cells":[]}
RESULT = {"organization":"Acme Ltd","organization_evidence":EVIDENCE,"document_type":"Report","warnings":[],"coverage":"complete","coverage_note":"All relevant source content reviewed.","facts":[{"organization":"Acme Ltd","organization_evidence":EVIDENCE,"label":"emissions","source_label":"emissions","value":"1234","unit":"tonnes","period":"2024","currency":None,"origin":"document_reported","evidence":EVIDENCE,"context_evidence":[]}]}


def response(result):
    return {"choices":[{"finish_reason":"stop","message":{"content":json.dumps(result)}}],"usage":{"total_tokens":100}}


def run_with(result, source=SOURCE):
    return {"result": project_result(FactResult.model_validate(result), source)}


def test_request_schema_and_header():
    def handler(request):
        assert request.headers["Authorization"] == "Bearer test-key"
        assert str(request.url) == "https://api.openai.com/v1/chat/completions"
        assert "test-key" not in str(request.url)
        body=json.loads(request.content)
        assert body["store"] is False
        assert body["response_format"]["type"] == "json_schema"
        assert body["response_format"]["json_schema"]["strict"] is True
        assert body["response_format"]["json_schema"]["schema"]["additionalProperties"] is False
        assert "maxItems" not in body["response_format"]["json_schema"]["schema"]["properties"]["facts"]
        assert "You are Harvey Specter" in body["messages"][0]["content"]
        indexed = {"document_type":"Report", "organization":None, "organization_ref":None,
                   "warnings":[], "coverage":"partial", "coverage_note":"Review", "facts":[]}
        return httpx.Response(200,json=response(indexed))
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        assert call_openai(SOURCE,"model","test-key",client)["result"]["facts"] == []


def test_rejects_fabrication_and_scaled_numbers():
    for key,value in [("page",9),("quote","invented 1,234"),("value","9999"),("value","1234 million"),("organization","Other")]:
        result=json.loads(json.dumps(RESULT))
        if key == "page": result["facts"][0]["evidence"]["unit_index"]=value
        elif key == "quote": result["facts"][0]["evidence"]["quote"]=value
        else: result["facts"][0][key]=value
        output = run_with(result)["result"]
        if key == "organization":
            assert output["facts"][0]["organization"] is None
        else:
            assert not output["facts"] and output["rejected_candidates"]


def test_cells_context_and_missing_organization():
    source={"version":1,"kind":"xlsx","summary":{"warnings":["cache missing"]},"units":[{"index":1,"label":"Sheet","cells":[{"coordinate":"A1","value":"emissions"},{"coordinate":"B1","value":1234},{"coordinate":"C1","value":"tonnes"}]}]}
    result=json.loads(json.dumps(RESULT)); result.update(organization=None,organization_evidence=None)
    result["facts"][0].update(organization=None,organization_evidence=None,period=None,evidence={"unit_index":1,"quote":"1234","cells":["B1"]},context_evidence=[{"unit_index":1,"quote":"emissions","cells":["A1"]},{"unit_index":1,"quote":"tonnes","cells":["C1"]}])
    output = run_with(result, source)["result"]
    assert output["coverage"] == "partial"
    assert any("Organization is missing" in warning for warning in output["warnings"])
    assert "cache missing" in output["warnings"]
    result["facts"][0]["evidence"]["cells"]=["Z99"]
    assert run_with(result,source)["result"]["rejected_candidates"]


def test_auth_network_and_incomplete():
    with httpx.Client(transport=httpx.MockTransport(lambda _:pytest.fail("provider called"))) as client:
        with pytest.raises(FactError): call_openai(SOURCE,"model","",client)
    for code,permanent in [(401,True),(429,False),(503,False)]:
        with httpx.Client(transport=httpx.MockTransport(lambda _,code=code:httpx.Response(code))) as client:
            with pytest.raises(FactError) as exc: call_openai(SOURCE,"model","key",client)
            assert exc.value.permanent is permanent
    with httpx.Client(transport=httpx.MockTransport(lambda _:httpx.Response(200,json={"choices":[{"finish_reason":"length","message":{}}]}))) as client:
        with pytest.raises(FactError): call_openai(SOURCE,"model","key",client)
    with httpx.Client(transport=httpx.MockTransport(lambda _: (_ for _ in ()).throw(httpx.ConnectError("offline")))) as client:
        with pytest.raises(FactError) as exc: call_openai(SOURCE,"model","key",client)
        assert not exc.value.permanent


def test_manual_queue_duplicate_and_fenced_completion(tmp_path):
    config=replace(Settings(),data_dir=tmp_path,openai_api_key="test-key",lease_seconds=10)
    with TestClient(create_app(config)) as client:
        doc=upload(client).json()["results"][0]["document"]
        assert process_one(Store(config))
        url=f'/api/documents/{doc["id"]}/facts'
        assert client.get(url).json()["status"] == "not_requested"
        assert client.post(url).json()["status"] == "queued"
        assert client.post(url).json()["status"] == "queued"
        store=Store(config); now=store.get(doc["id"])["updated_at"]+1
        first=fact_queue.claim(store,now=now); second=fact_queue.claim(store,now=now+11)
        assert not fact_queue.finish(store,first,output={"wrong":True},now=now+12)
        assert fact_queue.finish(store,second,output={"result":{"facts":[]}},now=now+12)
        assert client.post(url).json()["status"] == "ready"
        assert store.get(doc["id"])["status"] == "ready"


def test_missing_key_never_queues(tmp_path):
    config=replace(Settings(),data_dir=tmp_path,openai_api_key="")
    with TestClient(create_app(config)) as client:
        doc=upload(client).json()["results"][0]["document"]
        assert process_one(Store(config))
        assert not client.get("/api/facts/config").json()["configured"]
        url=f'/api/documents/{doc["id"]}/facts'
        assert client.post(url).status_code == 503
        assert client.get(url).json()["status"] == "not_requested"
        assert not process_fact(Store(config))


def test_numeric_evidence_does_not_match_substrings():
    result=json.loads(json.dumps(RESULT))
    result["facts"][0].update(value="123",unit=None,period=None,evidence={"unit_index":1,"quote":"123","cells":[]},context_evidence=[EVIDENCE])
    assert run_with(result)["result"]["rejected_candidates"]


def test_snapshot_is_frozen_and_temporary_retry_recovers(tmp_path):
    config=replace(Settings(),data_dir=tmp_path,openai_api_key="test-key",retry_base_seconds=0)
    with TestClient(create_app(config)) as client:
        doc=upload(client).json()["results"][0]["document"]
        store=Store(config)
        assert process_one(store)
        assert client.post(f'/api/documents/{doc["id"]}/facts').status_code == 202
        with store.connect() as db:
            db.execute("UPDATE source_units SET content_json=? WHERE document_id=?", ('{"index":1,"label":"changed","cells":[]}',doc["id"]))
        calls=[]
        def simulated(source, model, key, segment=None):
            calls.append(source)
            if len(calls)==1: raise FactError("temporary", permanent=False)
            return {"result":{"facts":[]},"model":model,"prompt_version":"reported-facts-v3","source_version":1}
        assert process_fact(store,simulated)
        assert fact_queue.get(store,doc["id"])["status"] == "queued"
        assert process_fact(store,simulated)
        result=fact_queue.get(store,doc["id"])
        assert result["status"] == "ready"
        assert calls[1]["units"][0]["label"] != "changed"
        assert len(result["output"]["segment_outputs"][0]["schema_hash"]) == 64
        assert len(result["output"]["segment_outputs"][0]["source_hash"]) == 64
        assert len(result["output"]["segment_outputs"][0]["document_hash"]) == 64


def test_numeric_quotes_cannot_match_inside_larger_source_number():
    for text, quote in [("emissions 142 tonnes", "42"), ("emissions 1,234 tonnes", "234"), ("emissions 1,234 tonnes", "234 tonnes")]:
        source={"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":text}]}
        result=json.loads(json.dumps(RESULT)); result.update(organization=None,organization_evidence=None)
        result["facts"][0].update(organization=None,organization_evidence=None,value=quote.split()[0],unit=None,period=None,evidence={"unit_index":1,"quote":quote,"cells":[]},context_evidence=[{"unit_index":1,"quote":text,"cells":[]}])
        assert run_with(result,source)["result"]["rejected_candidates"]
    source={"version":1,"kind":"xlsx","summary":{"warnings":[]},"units":[{"index":1,"label":"S","cells":[{"coordinate":"A1","value":142},{"coordinate":"B1","value":"emissions"}]}]}
    result=json.loads(json.dumps(RESULT)); result.update(organization=None,organization_evidence=None)
    result["facts"][0].update(organization=None,organization_evidence=None,value="42",unit=None,period=None,evidence={"unit_index":1,"quote":"42","cells":["A1"]},context_evidence=[{"unit_index":1,"quote":"emissions","cells":["B1"]}])
    assert run_with(result,source)["result"]["rejected_candidates"]


def test_schema_change_does_not_reuse_completed_result(tmp_path, monkeypatch):
    config = replace(Settings(), data_dir=tmp_path, openai_api_key="test-key")
    with TestClient(create_app(config)) as client:
        doc = upload(client).json()["results"][0]["document"]
        store = Store(config)
        assert process_one(store)
        assert fact_queue.enqueue(store, doc["id"]) == "queued"
        job = fact_queue.claim(store)
        assert fact_queue.finish(store, job, output={"old": True})
        assert fact_queue.enqueue(store, doc["id"]) == "existing"
        monkeypatch.setattr(fact_queue, "SCHEMA_HASH", "changed-schema")
        assert fact_queue.enqueue(store, doc["id"]) == "queued"


def test_partial_coverage_and_fact_cap_are_explicit():
    result = json.loads(json.dumps(RESULT))
    result["coverage"] = "partial"
    result["coverage_note"] = "More figures could be relevant."
    output = run_with(result)["result"]
    assert output["coverage"] == "partial"
    assert any("incomplete" in warning for warning in output["warnings"])
    result["facts"] *= 61
    with pytest.raises(ValueError): run_with(result)


@pytest.mark.parametrize("field", ["label", "unit", "period", "currency"])
def test_unsupported_context_names_fact_and_field(field):
    result = json.loads(json.dumps(RESULT))
    result["facts"][0][field] = "unsupported-context"
    output = run_with(result)["result"]
    if field == "label":
        assert output["facts"][0]["label"] == "unsupported-context"  # AI-proposed display label
    else:
        assert output["facts"][0][field] is None
        assert any(issue["field"] == field for issue in output["review_issues"])


def test_mixed_candidates_preserve_supported_and_audit_rejected():
    result = json.loads(json.dumps(RESULT))
    good = result["facts"][0]
    bad = json.loads(json.dumps(good))
    bad["value"] = "9999"
    result["facts"].append(bad)
    output = run_with(result)["result"]
    assert len(output["facts"]) == 1
    assert output["facts"][0]["value"] == "1234"
    assert output["rejected_candidates"][0]["candidate"]["value"] == "9999"
    assert len(output["proposed_facts"]) == 2
    assert output["coverage"] == "partial"


def test_metadata_is_case_sensitive_and_never_joined_across_quotes():
    source = {"version":1,"kind":"xlsx","summary":{"warnings":[]},"units":[{"index":1,"label":"S","cells":[{"coordinate":"A1","value":"Power"},{"coordinate":"B1","value":42},{"coordinate":"C1","value":"mWh"},{"coordinate":"D1","value":"kg"},{"coordinate":"E1","value":"CO2"}]}]}
    result = json.loads(json.dumps(RESULT)); result.update(organization="Unknown Inc",organization_evidence=EVIDENCE)
    fact = result["facts"][0]
    fact.update(organization="Unknown Inc",organization_evidence=EVIDENCE,label="Power use",source_label="Power",value="42",unit="MWh",period="kg CO2",currency=None,evidence={"unit_index":1,"quote":"42","cells":["B1"]},context_evidence=[{"unit_index":1,"quote":value,"cells":[cell]} for cell,value in [("A1","Power"),("C1","mWh"),("D1","kg"),("E1","CO2")]])
    output=run_with(result,source)["result"]
    assert output["organization"] is None and output["facts"][0]["organization"] is None
    assert output["facts"][0]["unit"] is None and output["facts"][0]["period"] is None
    assert output["facts"][0]["value"] == "42"
    assert {issue["field"] for issue in output["review_issues"]} >= {"unit","period","organization"}
    assert output["coverage"] == "partial"


def test_revalidation_uses_snapshot_without_key_and_preserves_rejected_audit(tmp_path):
    config = replace(Settings(), data_dir=tmp_path, openai_api_key="test-key")
    with TestClient(create_app(config)) as client:
        doc=upload(client).json()["results"][0]["document"]
        store=Store(config)
        assert process_one(store)
        url=f'/api/documents/{doc["id"]}/facts'
        assert client.post(url).status_code == 202
        bad = json.loads(json.dumps(RESULT))
        bad.update(organization=None, organization_evidence=None)
        bad["facts"][0].update(organization=None, organization_evidence=None, source_label="Scope 1", value="999", evidence={"unit_index":1,"quote":"Scope 1","cells":["A1"]})
        def caller(source, model, key, segment=None):
            return {"result": run_with(bad,source)["result"], "proposed_result": bad, "index_version": "legacy-v3", "validator_version": "facts-validation-v4"}
        assert process_fact(store,caller)
        first=client.get(url).json()
        assert first["status"] == "failed"
        assert first["output"]["result"]["rejected_candidates"][0]["candidate"]["value"] == "999"
        with store.connect() as db:
            db.execute("UPDATE source_units SET content_json=? WHERE document_id=?", ('{"index":1,"label":"changed","cells":[]}',doc["id"]))
    no_key = replace(config, openai_api_key="")
    with TestClient(create_app(no_key)) as client:
        replay=client.post(url + '/revalidate')
        assert replay.status_code == 409  # v3 candidates require their original validator


def test_revalidation_unavailable_for_old_or_active_jobs(tmp_path):
    config=replace(Settings(),data_dir=tmp_path,openai_api_key="test-key")
    with TestClient(create_app(config)) as client:
        doc=upload(client).json()["results"][0]["document"]
        store=Store(config); assert process_one(store)
        url=f'/api/documents/{doc["id"]}/facts'
        assert client.post(url + '/revalidate').status_code == 409
        assert client.post(url).status_code == 202
        assert client.post(url + '/revalidate').status_code == 409
        job=fact_queue.claim(store)
        assert fact_queue.finish(store,job,output={"old":"no saved candidates"})
        assert client.post(url + '/revalidate').status_code == 409


def test_reported_symbols_are_preserved_but_codes_are_not_inferred():
    source = {"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":"Share 34.0% and cost $100."}]}
    result=json.loads(json.dumps(RESULT)); result.update(organization=None,organization_evidence=None)
    fact=result["facts"][0]
    fact.update(organization=None,organization_evidence=None,label="Share",source_label="Share",value="34.0",unit="%",period=None,currency="$",evidence={"unit_index":1,"quote":"Share 34.0% and cost $100.","cells":[]})
    output=run_with(result,source)["result"]
    assert output["facts"][0]["unit"] == "%" and output["facts"][0]["currency"] == "$"
    fact["currency"]="USD"
    output=run_with(result,source)["result"]
    assert output["facts"][0]["currency"] is None
    assert any(issue["field"] == "currency" for issue in output["review_issues"])


def test_optional_context_rejection_preserves_supported_bill_value():
    source = {"version": 1, "kind": "pdf", "summary": {"warnings": []},
              "units": [{"index": 1, "label": "Page 1", "text": "Energy consumption 1,234 kWh. Billing period October 2024."}]}
    result = json.loads(json.dumps(RESULT))
    result.update(organization=None, organization_evidence=None)
    fact = result["facts"][0]
    fact.update(organization=None, organization_evidence=None, label="Electricity used", source_label="Energy consumption",
                value="1234", unit="kWh", period="November 2024", currency=None,
                evidence={"unit_index": 1, "quote": "Energy consumption 1,234 kWh", "cells": []},
                context_evidence=[{"unit_index": 1, "quote": "Billing period October 2024", "cells": []}])
    output = run_with(result, source)["result"]
    assert len(output["facts"]) == 1
    assert output["facts"][0]["value"] == "1234"
    assert output["facts"][0]["unit"] == "kWh"
    assert output["facts"][0]["period"] is None
    assert output["review_issues"][0]["proposed"] == "November 2024"
    assert output["coverage"] == "partial"


@pytest.mark.parametrize("quote", ["Energy <42 kWh", "Energy ≥42 kWh", "Energy approximately 42 kWh", "Energy 42 million kWh"])
def test_numeric_qualifier_or_scale_cannot_be_dropped(quote):
    source = {"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":quote}]}
    result=json.loads(json.dumps(RESULT)); result.update(organization=None,organization_evidence=None)
    result["facts"][0].update(organization=None,organization_evidence=None,source_label="Energy",value="42",unit="kWh",period=None,currency=None,evidence={"unit_index":1,"quote":quote,"cells":[]},context_evidence=[])
    output=run_with(result,source)["result"]
    assert output["facts"] == []
    assert "qualifier or scale" in output["rejected_candidates"][0]["reason"]


def test_locates_unique_label_on_same_page_with_review_warning():
    source={"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":"Current Month Bill Previous Outstanding APR-24 47407 632023.17"}]}
    result=json.loads(json.dumps(RESULT)); result.update(organization=None,organization_evidence=None)
    result["facts"][0].update(organization=None,organization_evidence=None, label="Current bill",source_label="Current Month Bill",value="632023.17",unit=None,period=None,currency=None,evidence={"unit_index":1,"quote":"632023.17","cells":[]},context_evidence=[])
    output=run_with(result,source)["result"]
    assert output["facts"] == []
    assert len(output["review_candidates"]) == 1
    assert output["review_candidates"][0]["context_evidence"][0] == {"unit_index":1,"quote":"Current Month Bill","cells":[],"located_locally":True}
    assert output["review_issues"][0]["field"] == "source_label"
    assert "inferred" in output["review_issues"][0]["reason"]
    assert output["coverage"] == "partial"
    source["units"][0]["text"] += " Current Month Bill"
    assert run_with(result,source)["result"]["rejected_candidates"]
    source["units"][0]["text"] = "Current Month Bill 632023.17"
    result["facts"][0]["context_evidence"] = [{"unit_index":1,"quote":"invented context","cells":[]}]
    assert run_with(result,source)["result"]["rejected_candidates"]


def test_indexed_source_refs_validate_and_do_not_invent_values():
    from app.facts import IndexedFactResult, project_indexed
    source={"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":"Current Month Bill\nAPR-24 47407 632023.17"}]}
    fact={"label":"Current bill","source_label":"Current Month Bill","label_ref":"p1:l1","value":"632023.17","value_ref":"p1:l2","organization":None,"organization_ref":None,"unit":None,"unit_ref":None,"period":"APR-24","period_ref":"p1:l2","currency":None,"currency_ref":None,"origin":"document_reported"}
    data={"document_type":"Bill","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    output=project_indexed(IndexedFactResult.model_validate(data),source)
    assert len(output["facts"]) == 1 and not output["review_candidates"]
    assert output["facts"][0]["evidence"]["quote"] == "APR-24 47407 632023.17"
    assert output["facts"][0]["period"] == "APR-24"
    assert output["facts"][0]["association"] == "same_unit_supported"
    fact["value_ref"]="p1:l99"
    assert project_indexed(IndexedFactResult.model_validate(data),source)["rejected_candidates"]
    fact["value_ref"]="p1:l2"; fact["value"]="999"
    assert project_indexed(IndexedFactResult.model_validate(data),source)["rejected_candidates"]


def test_indexed_spreadsheet_missing_cache_is_rejected():
    from app.facts import IndexedFactResult, project_indexed
    source={"version":1,"kind":"xlsx","summary":{"warnings":["cache missing"]},"units":[{"index":1,"label":"S","cells":[{"coordinate":"A1","type":"s","value":"Emissions"},{"coordinate":"B1","type":"f","value":None,"formula":"=A2","cached_value":None}]}]}
    fact={"label":"Emissions","source_label":"Emissions","label_ref":"s1:A1","value":"42","value_ref":"s1:B1","organization":None,"organization_ref":None,"unit":None,"unit_ref":None,"period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
    data={"document_type":"Report","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    output=project_indexed(IndexedFactResult.model_validate(data),source)
    assert output["rejected_candidates"] and not output["facts"]


def test_prior_run_archived_before_replacement(tmp_path, monkeypatch):
    config=replace(Settings(),data_dir=tmp_path,openai_api_key="test-key")
    with TestClient(create_app(config)) as client:
        doc=upload(client).json()["results"][0]["document"]
        store=Store(config); assert process_one(store)
        assert fact_queue.enqueue(store,doc["id"]) == "queued"
        job=fact_queue.claim(store)
        assert fact_queue.finish(store,job,output={"result":{"facts":[]},"proposed_result":{"legacy":True}})
        old_snapshot=job["source_snapshot"]
        monkeypatch.setattr(fact_queue,"SCHEMA_HASH","different-schema")
        assert fact_queue.enqueue(store,doc["id"]) == "queued"
        history=client.get(f'/api/documents/{doc["id"]}/facts/history').json()["runs"]
        assert len(history)==1 and history[0]["job"]["source_snapshot"] == old_snapshot
        assert json.loads(history[0]["job"]["result_json"])["segment_outputs"][0]["proposed_result"]["legacy"] is True


def test_sheet_title_index_can_support_period_but_missing_cache_cannot_support_value():
    from app.evidence_index import build_index
    from app.facts import IndexedFactResult, project_indexed
    source={"version":1,"kind":"xlsx","summary":{"warnings":[]},"units":[{"index":1,"label":"Monthly Summary 2026","cells":[{"coordinate":"A1","type":"s","value":"Energy"},{"coordinate":"B1","type":"n","value":42}]}]}
    assert build_index(source)["records"][0]["id"] == "s1:title"
    fact={"label":"Energy","source_label":"Energy","label_ref":"s1:A1","value":"42","value_ref":"s1:B1","organization":None,"organization_ref":None,"unit":None,"unit_ref":None,"period":"2026","period_ref":"s1:title","currency":None,"currency_ref":None,"origin":"document_reported"}
    data={"document_type":"Report","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    output=project_indexed(IndexedFactResult.model_validate(data),source)
    assert output["facts"][0]["period"] == "2026"
    assert output["facts"][0]["association"] == "same_unit_supported"
    fact["period"]="Monthly Summary 2026"
    assert project_indexed(IndexedFactResult.model_validate(data),source)["facts"][0]["period"] == "Monthly Summary 2026"


def test_same_line_supported_fact_automatically_checks_but_remains_model_proposed():
    from app.facts import IndexedFactResult, project_indexed
    source={"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":"Energy used 42 kWh"}]}
    fact={"label":"Energy used","source_label":"Energy used","label_ref":"p1:l1","value":"42","value_ref":"p1:l1","organization":None,"organization_ref":None,"unit":"kWh","unit_ref":"p1:l1","period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
    data={"document_type":"Report","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    output=project_indexed(IndexedFactResult.model_validate(data),source)
    assert output["facts"][0]["value"] == "42" and not output["review_candidates"]


def test_unknown_optional_reference_is_cleared_and_flagged():
    from app.facts import IndexedFactResult, project_indexed
    source={"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":"Energy used 42 kWh"}]}
    fact={"label":"Energy","source_label":"Energy used","label_ref":"p1:l1","value":"42","value_ref":"p1:l1","organization":None,"organization_ref":None,"unit":"kWh","unit_ref":"p1:l99","period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
    data={"document_type":"Report","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    output=project_indexed(IndexedFactResult.model_validate(data),source)
    assert output["facts"][0]["unit"] is None
    assert any(issue["field"] == "unit" for issue in output["review_issues"])


def test_segment_manifest_covers_all_cells_once_and_context_is_separate():
    from app.segments import build_manifest
    cells=[{"coordinate":f"A{i}","type":"n","value":i} for i in range(1,14)]
    source={"version":1,"kind":"xlsx","summary":{"warnings":[]},"units":[{"index":1,"label":"2026","cells":cells}]}
    segments=build_manifest(source)
    target_ids=[item for segment in segments for item in segment["target_ids"]]
    assert len(segments)==4 and len(target_ids)==len(set(target_ids))==13
    assert set(target_ids)=={f"s1:A{i}" for i in range(1,14)}
    assert "s1:A1" in {r["id"] for r in segments[1]["records"]}
    assert "s1:A1" not in segments[1]["target_ids"]
    assert segments[0]["extraction_hints"]["prefer_numeric_targets_under_unit_headers"] is True


def test_xlsx_targets_prefer_numeric_cells_and_keep_headers_in_context():
    from app.segments import build_manifest
    cells=[
        {"coordinate":"A1","type":"s","value":"Month"},
        {"coordinate":"B1","type":"s","value":"Grid Supply (kWh)"},
        {"coordinate":"A2","type":"s","value":"January"},
        {"coordinate":"B2","type":"n","value":100},
        {"coordinate":"A3","type":"s","value":"February"},
        {"coordinate":"B3","type":"n","value":200},
    ]
    source={"version":1,"kind":"xlsx","summary":{"warnings":[]},"units":[{"index":1,"label":"Monthly","cells":cells}]}
    segments=build_manifest(source)
    target_ids=[item for segment in segments for item in segment["target_ids"]]
    assert set(target_ids)=={"s1:B2","s1:B3"}
    records={r["id"] for segment in segments for r in segment["records"]}
    assert {"s1:title","s1:A1","s1:B1","s1:A2","s1:B2","s1:A3","s1:B3"} <= records
    assert "s1:B1" in {r["id"] for r in segments[0]["records"] if r["id"] not in segments[0]["target_ids"]}


def test_pdf_spans_require_contiguous_same_page_and_segment_targets():
    from app.segments import build_manifest
    from app.facts import IndexedFactResult, project_indexed
    source={"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":"Quality commitment\nDeliver safe\nproducts"}]}
    segment=build_manifest(source)[0]
    fact={"label":"Quality","source_label":"Quality commitment","label_ref":"p1:l1","label_refs":[],"value":"Deliver safe products","value_ref":"p1:l2","value_refs":["p1:l2","p1:l3"],"organization":None,"organization_ref":None,"unit":None,"unit_ref":None,"period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
    data={"document_type":"Policy","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    output=project_indexed(IndexedFactResult.model_validate(data),source,segment)
    assert output["facts"][0]["evidence"]["quote"] == "Deliver safe\nproducts"
    assert output["facts"][0]["association"] == "same_unit_supported"
    for refs in (["p1:l2","p1:l1"],["p1:l2","p2:l3"],["p1:l2","p1:l4"]):
        fact["value_refs"]=refs
        assert project_indexed(IndexedFactResult.model_validate(data),source,segment)["rejected_candidates"]
    fact["value_refs"]=["p1:l2","p1:l3"]
    segment["target_ids"]=["p1:l1"]
    assert project_indexed(IndexedFactResult.model_validate(data),source,segment)["rejected_candidates"]


def test_segmented_worker_resumes_completed_and_fences_stale(tmp_path):
    book=__import__("openpyxl").Workbook()
    for i in range(1,10): book.active[f"A{i}"]=i
    import io
    payload=io.BytesIO(); book.save(payload)
    config=replace(Settings(),data_dir=tmp_path,openai_api_key="test-key",lease_seconds=10,retry_base_seconds=0)
    with TestClient(create_app(config)) as client:
        doc=upload(client,payload.getvalue()).json()["results"][0]["document"]
        store=Store(config); assert process_one(store)
        assert client.post(f'/api/documents/{doc["id"]}/facts').status_code==202
        first=fact_queue.claim(store); assert first
        assert fact_queue.finish(store,first,output={"result":{"facts":[],"review_candidates":[],"warnings":[]}})
        second=fact_queue.claim(store); assert second["segment_id"] != first["segment_id"]
        abandoned_time = __import__("time").time()
        recovered=fact_queue.claim(store,now=abandoned_time+11)
        assert recovered["segment_id"]==second["segment_id"]
        assert not fact_queue.finish(store,second,output={"result":{"facts":[]}},now=abandoned_time+12)
        assert fact_queue.finish(store,recovered,output={"result":{"facts":[],"review_candidates":[],"warnings":[]}},now=abandoned_time+12)
        third=fact_queue.claim(store,now=abandoned_time+12)
        assert third["segment_id"] not in (first["segment_id"],second["segment_id"])
        assert fact_queue.finish(store,third,output={"result":{"facts":[],"review_candidates":[],"warnings":[]}},now=abandoned_time+13)
        output=fact_queue.get(store,doc["id"])["output"]
        assert output["completed_segments"]==3 and output["total_target_records"]==9
        assert output["operational_coverage"]=="complete"


def test_failed_segment_retry_preserves_completed_segment(tmp_path):
    book=__import__("openpyxl").Workbook()
    for i in range(1,6): book.active[f"A{i}"]=i
    import io
    payload=io.BytesIO(); book.save(payload)
    config=replace(Settings(),data_dir=tmp_path,openai_api_key="test-key")
    with TestClient(create_app(config)) as client:
        doc=upload(client,payload.getvalue()).json()["results"][0]["document"]
        store=Store(config); assert process_one(store)
        assert fact_queue.enqueue(store,doc["id"]) == "queued"
        one=fact_queue.claim(store); assert fact_queue.finish(store,one,output={"result":{"facts":[]}})
        two=fact_queue.claim(store); assert fact_queue.finish(store,two,error="unavailable",permanent=True)
        assert fact_queue.get(store,doc["id"])["output"]["operational_coverage"] == "partial"
        assert fact_queue.enqueue(store,doc["id"]) == "queued"
        retry=fact_queue.claim(store)
        assert retry["segment_id"] == two["segment_id"]
        assert fact_queue.finish(store,retry,output={"result":{"facts":[]}})
        status=fact_queue.get(store,doc["id"])
        assert status["output"]["completed_segments"] == 2
        assert status["segments"][0]["attempts"] == 1


def test_conflicting_segment_results_are_not_silently_accepted(tmp_path):
    book=__import__("openpyxl").Workbook()
    for i in range(1,6): book.active[f"A{i}"]=i
    import io
    payload=io.BytesIO(); book.save(payload)
    config=replace(Settings(),data_dir=tmp_path,openai_api_key="test-key")
    with TestClient(create_app(config)) as client:
        doc=upload(client,payload.getvalue()).json()["results"][0]["document"]
        store=Store(config); assert process_one(store); assert fact_queue.enqueue(store,doc["id"]) == "queued"
        base={"source_label":"Energy","value":"42","unit":"kWh","period":None,"currency":None,"organization":"A","evidence":{"reference_id":"s1:A1","reference_ids":["s1:A1"]}}
        first=fact_queue.claim(store); assert fact_queue.finish(store,first,output={"result":{"facts":[base]}})
        other={**base,"organization":"B"}
        second=fact_queue.claim(store); assert fact_queue.finish(store,second,output={"result":{"facts":[other]}})
        result=fact_queue.get(store,doc["id"])["output"]["result"]
        assert not result["facts"] and len(result["review_candidates"]) == 2
        assert any("Conflicting" in warning for warning in result["warnings"])


def test_segmented_revalidation_uses_snapshot_and_no_provider(tmp_path):
    config=replace(Settings(),data_dir=tmp_path,openai_api_key="test-key")
    with TestClient(create_app(config)) as client:
        doc=upload(client).json()["results"][0]["document"]
        store=Store(config); assert process_one(store)
        assert fact_queue.enqueue(store,doc["id"]) == "queued"
        job=fact_queue.claim(store)
        segment=json.loads(job["segment_payload"])
        refs=segment["target_ids"]
        fact={"label":"Scope","source_label":"Scope 1","label_ref":refs[0],"value":"Scope 1","value_ref":refs[0],"organization":None,"organization_ref":None,"unit":None,"unit_ref":None,"period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
        candidate={"document_type":"Workbook","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
        from app.facts import IndexedFactResult, project_indexed
        source=json.loads(job["source_snapshot"])
        output={"index_version":"source-index-v1","proposed_result":candidate,"result":project_indexed(IndexedFactResult.model_validate(candidate),source,segment),"model":"test"}
        assert fact_queue.finish(store,job,output=output)
        with store.connect() as db:
            db.execute("UPDATE source_units SET content_json=? WHERE document_id=?", ('{"index":1,"label":"changed","cells":[]}',doc["id"]))
        assert client.post(f'/api/documents/{doc["id"]}/facts/revalidate').status_code == 200
        result=client.get(f'/api/documents/{doc["id"]}/facts').json()["output"]
        assert result["result"]["facts"][0]["value"] == "Scope 1"


def test_narrative_column_detector_preserves_raw_and_ignores_numeric_tables():
    from app.extraction import narrative_columns
    class Crop:
        def __init__(self, label): self.label=label
        def extract_text(self): return self.label
    class Page:
        width=600; height=800
        def __init__(self, numeric=False): self.numeric=numeric
        def extract_words(self):
            return [{"x0":x,"x1":x+20,"top":100+i,"bottom":105+i,"text":str(i) if self.numeric else "word"} for i in range(80) for x in (100,400)]
        def crop(self,bbox): return Crop("column content")
    blocks=narrative_columns(Page())
    assert blocks and any(b["label"] == "Left column" for b in blocks) and any(b["label"] == "Right column" for b in blocks)
    assert narrative_columns(Page(numeric=True)) is None


def test_indexed_pdf_block_spans_stay_in_same_block():
    from app.evidence_index import build_index
    from app.facts import IndexedFactResult, project_indexed
    from app.segments import build_manifest
    source={"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":"raw interleaved text","reading_blocks":[{"label":"Left column","bbox":[0,0,100,100],"text":"Quality\nSafe food\nfor customers"},{"label":"Right column","bbox":[100,0,200,100],"text":"Other"}]}]}
    ids=[r["id"] for r in build_index(source)["records"]]
    assert ids == ["p1:b1:l1","p1:b1:l2","p1:b1:l3","p1:b2:l1"]
    fact={"label":"Quality","source_label":"Quality","label_ref":"p1:b1:l1","value":"Safe food for customers","value_ref":"p1:b1:l2","value_refs":["p1:b1:l2","p1:b1:l3"],"organization":None,"organization_ref":None,"unit":None,"unit_ref":None,"period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
    data={"document_type":"Policy","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    segment=build_manifest(source)[0]
    output=project_indexed(IndexedFactResult.model_validate(data),source,segment)
    assert output["facts"][0]["evidence"]["quote"] == "Safe food\nfor customers"
    assert output["facts"][0]["association"] == "same_unit_supported"
    fact["value_refs"]=["p1:b1:l2","p1:b2:l1"]
    assert project_indexed(IndexedFactResult.model_validate(data),source,segment)["rejected_candidates"]


def test_literal_value_metadata_and_formula_cache_flags():
    from app.evidence_index import build_index
    source={"version":1,"kind":"xlsx","summary":{"warnings":[]},"units":[{"index":1,"label":"S","cells":[{"coordinate":"A1","type":"n","value":200823},{"coordinate":"A2","type":"f","value":None,"formula":"=1+2","cached_value":None},{"coordinate":"A3","type":"f","value":3,"formula":"=1+2","cached_value":3}]}]}
    records={record["id"]:record for record in build_index(source)["records"]}
    assert records["s1:A1"]["value_origin"] == "literal" and records["s1:A1"]["typed_value"] == 200823
    assert records["s1:A1"]["value_available"] and "cached_value" not in records["s1:A1"]
    assert records["s1:A2"]["value_origin"] == "missing_formula_cache" and not records["s1:A2"]["value_available"]
    assert records["s1:A3"]["value_origin"] == "formula_cache" and records["s1:A3"]["typed_value"] == 3


def test_local_table_review_only_for_unproposed_target_values_with_unit_headers():
    from app.segments import build_manifest
    from app.facts import IndexedFactResult, project_indexed
    cells=[{"coordinate":"A2","type":"s","value":"Month"},{"coordinate":"F2","type":"s","value":"Grid Supply (kWh)"},{"coordinate":"G2","type":"s","value":"Remarks"},{"coordinate":"A3","type":"s","value":"January"},{"coordinate":"F3","type":"n","value":200823},{"coordinate":"G3","type":"n","value":10},{"coordinate":"F4","type":"f","value":None,"formula":"=1+1","cached_value":None}]
    source={"version":1,"kind":"xlsx","summary":{"warnings":[]},"units":[{"index":1,"label":"S","cells":cells}]}
    segment=build_manifest(source)[0]
    data={"document_type":"Workbook","organization":None,"organization_ref":None,"facts":[],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    result=project_indexed(IndexedFactResult.model_validate(data),source,segment)
    assert result["review_candidates"] == []
    checked=[f for f in result["facts"] if f.get("method") == "deterministic_table_fact"]
    assert len(checked) == 1 and checked[0]["value"] == "200823"
    assert checked[0]["source_label"] == "Grid Supply (kWh)" and checked[0]["period"] == "January"
    assert checked[0]["evidence"]["reference_id"] == "s1:F3" and checked[0]["context_evidence"][0]["reference_id"] == "s1:F2"
    assert checked[0]["context_evidence"][1]["quote"] == "January"
    assert checked[0]["association"] == "column_header_supported"
    # Numeric remarks without a unit-bearing header stay out of Checked.
    assert all(f["evidence"]["reference_id"] != "s1:G3" for f in result["facts"])
    # String headers are context, not AI targets.
    assert "s1:F2" not in segment["target_ids"] and "s1:F3" in segment["target_ids"]
    # Rejected model proposals do not block structural header→value Checked facts.
    bad={"label":"Wrong","source_label":"Invented","label_ref":"s1:F2","value":"999","value_ref":"s1:F3","organization":None,"organization_ref":None,"unit":None,"unit_ref":None,"period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
    data["facts"]=[bad]
    result=project_indexed(IndexedFactResult.model_validate(data),source,segment)
    assert result["rejected_candidates"]
    assert any(f.get("method") == "deterministic_table_fact" and f["value"] == "200823" for f in result["facts"])
    # Already-Checked model facts for the same value_ref are not duplicated.
    good={"label":"Grid Supply (kWh)","source_label":"Grid Supply (kWh)","label_ref":"s1:F2","value":"200823","value_ref":"s1:F3","organization":None,"organization_ref":None,"unit":None,"unit_ref":None,"period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
    data["facts"]=[good]
    result=project_indexed(IndexedFactResult.model_validate(data),source,segment)
    mapped=[f for f in result["facts"] if f.get("method") == "deterministic_table_fact"]
    assert result["facts"] and not mapped
    data["facts"]=[]; segment["target_ids"].remove("s1:F3")
    assert project_indexed(IndexedFactResult.model_validate(data),source,segment)["facts"] == []


def test_local_revalidation_adds_legacy_table_candidates_without_provider(tmp_path):
    import io
    import openpyxl
    book=openpyxl.Workbook(); sheet=book.active
    sheet["A2"]="Month"; sheet["F2"]="Grid Supply (kWh)"; sheet["A3"]="January"; sheet["F3"]=200823
    payload=io.BytesIO(); book.save(payload)
    config=replace(Settings(),data_dir=tmp_path,openai_api_key="test-key")
    with TestClient(create_app(config)) as client:
        doc=upload(client,payload.getvalue()).json()["results"][0]["document"]
        store=Store(config); assert process_one(store)
        assert fact_queue.enqueue(store,doc["id"]) == "queued"
        job=fact_queue.claim(store)
        proposed={"document_type":"Workbook","organization":None,"organization_ref":None,"facts":[],"warnings":[],"coverage":"partial","coverage_note":"No numeric values found."}
        legacy_result={"document_type":"Workbook","organization":None,"facts":[],"review_candidates":[],"rejected_candidates":[],"review_issues":[],"warnings":[],"coverage":"partial","coverage_note":"No numeric values found."}
        output={"index_version":"source-index-v1","proposed_result":proposed,"result":legacy_result,"model":"test"}
        assert fact_queue.finish(store,job,output=output)
        assert client.post(f'/api/documents/{doc["id"]}/facts/revalidate').status_code == 200
        saved=client.get(f'/api/documents/{doc["id"]}/facts').json()["output"]
        found=[c for c in saved["result"]["facts"] if c["evidence"]["reference_id"] == "s1:F3"]
        assert len(found)==1 and found[0]["value"]=="200823"
        assert found[0]["method"] == "deterministic_table_fact"
        assert found[0]["period"] == "January"
        assert not saved["result"]["review_candidates"] and saved["segment_outputs"][0]["proposed_result"] == proposed
        from app.facts import VALIDATOR_VERSION
        assert saved["segment_outputs"][0]["validator_version"] == VALIDATOR_VERSION


def test_openai_strict_schema_all_object_fields_required():
    from app.facts import provider_schema
    schema = provider_schema()
    def check(node):
        if isinstance(node, dict):
            if node.get("type") == "object":
                assert node.get("additionalProperties") is False
                assert set(node.get("properties", {})) == set(node.get("required", []))
            for value in node.values(): check(value)
        elif isinstance(node, list):
            for value in node: check(value)
    check(schema)


def test_openai_refusal_truncation_quota_and_no_secret_errors():
    cases=[({"choices":[{"finish_reason":"stop","message":{"refusal":"private refusal"}}]},"refused"),
           ({"choices":[{"finish_reason":"length","message":{"content":"{"}}]},"complete"),
           ({"choices":[{"finish_reason":"stop","message":{"content":"not json"}}]},"invalid")]
    for payload, phrase in cases:
        with httpx.Client(transport=httpx.MockTransport(lambda _, payload=payload: httpx.Response(200,json=payload))) as client:
            with pytest.raises(FactError) as exc: call_openai(SOURCE,"gpt-4o-mini","secret-key",client)
            assert phrase in str(exc.value).lower() and "secret-key" not in str(exc.value)
    with httpx.Client(transport=httpx.MockTransport(lambda _:httpx.Response(429,json={"error":{"code":"insufficient_quota"}}))) as client:
        with pytest.raises(FactError) as exc: call_openai(SOURCE,"gpt-4o-mini","secret-key",client)
        assert exc.value.permanent and "quota" in str(exc.value).lower()


def test_old_provider_jobs_never_sent_and_archive_includes_segments(tmp_path):
    config=replace(Settings(),data_dir=tmp_path,openai_api_key="test-key")
    with TestClient(create_app(config)) as client:
        doc=upload(client).json()["results"][0]["document"]
        store=Store(config); assert process_one(store)
        assert fact_queue.enqueue(store,doc["id"]) == "queued"
        with store.connect() as db:
            db.execute("UPDATE fact_jobs SET provider='gemini',model='gemini-3.5-flash-lite' WHERE document_id=?",(doc["id"],))
        def never_call(*args, **kwargs): pytest.fail("old job was sent to OpenAI")
        assert not process_fact(store,never_call)
        prior=fact_queue.get(store,doc["id"])
        assert prior["status"] == "failed" and prior["provider"] == "gemini"
        assert fact_queue.enqueue(store,doc["id"]) == "queued"
        history=client.get(f'/api/documents/{doc["id"]}/facts/history').json()["runs"]
        assert history[0]["job"]["provider"] == "gemini"
        assert history[0]["job"]["segments"] and history[0]["job"]["segments"][0]["payload_json"]
        assert fact_queue.get(store,doc["id"])["provider"] == "openai"


def test_recover_exact_header_reference_is_review_only_and_audited():
    from app.facts import IndexedFactResult, project_indexed
    source={"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":"Current Month Bill  Units Consumed\nAPR-24 47407 632023.17"}]}
    segment={"id":"p1-1","kind":"pdf","target_ids":["p1:l1","p1:l2"],"records":[{"id":"p1:l1","unit_index":1,"text":"Current Month Bill  Units Consumed","cells":[]},{"id":"p1:l2","unit_index":1,"text":"APR-24 47407 632023.17","cells":[]}]}
    fact={"label":"Current bill","source_label":"Current Month Bill","label_ref":"p1:l2","value":"632023.17","value_ref":"p1:l2","organization":None,"organization_ref":None,"unit":None,"unit_ref":None,"period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
    data={"document_type":"Bill","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    result=project_indexed(IndexedFactResult.model_validate(data),source,segment)
    assert not result["facts"] and result["review_candidates"][0]["method"] == "local_exact_reference_recovery"
    item=result["review_candidates"][0]
    assert item["fact_index"] == 1 and item["recovery_audit"][0]["original_ids"] == ["p1:l2"]
    assert item["recovery_audit"][0]["recovered_ids"] == ["p1:l1"]
    assert result["proposed_facts"][0]["label_ref"] == "p1:l2"
    # Duplicate exact matches are ambiguous, not a reason to choose one.
    segment["records"].append({"id":"p1:l3","unit_index":1,"text":"Current Month Bill","cells":[]})
    source["units"][0]["text"] += "\nCurrent Month Bill"
    assert project_indexed(IndexedFactResult.model_validate(data),source,segment)["rejected_candidates"]


def test_recovery_does_not_cross_page_and_flags_optional_metadata():
    from app.facts import IndexedFactResult, project_indexed
    source={"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":"Value 42 kWh"},{"index":2,"label":"Page 2","text":"Energy"}]}
    segment={"id":"p1","kind":"pdf","target_ids":["p1:l1"],"records":[{"id":"p1:l1","unit_index":1,"text":"Value 42 kWh","cells":[]},{"id":"p2:l1","unit_index":2,"text":"Energy","cells":[]}]}
    fact={"label":"Energy","source_label":"Energy","label_ref":"p1:l1","value":"42","value_ref":"p1:l1","organization":None,"organization_ref":None,"unit":"kWh","unit_ref":"p1:l999","period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
    data={"document_type":"Report","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    assert project_indexed(IndexedFactResult.model_validate(data),source,segment)["rejected_candidates"]
    fact["source_label"]="Value"
    output=project_indexed(IndexedFactResult.model_validate(data),source,segment)
    assert output["facts"][0]["unit"] is None and any(i["field"]=="unit" for i in output["review_issues"])


def test_contact_details_are_rejected_explicitly():
    from app.facts import IndexedFactResult, project_indexed
    source={"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":"Phone 1234567890"}]}
    fact={"label":"Phone","source_label":"Phone","label_ref":"p1:l1","value":"1234567890","value_ref":"p1:l1","organization":None,"organization_ref":None,"unit":None,"unit_ref":None,"period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
    data={"document_type":"Bill","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    assert "outside" in project_indexed(IndexedFactResult.model_validate(data),source)["rejected_candidates"][0]["reason"]


def test_duplicate_span_references_recovered_and_policy_email_not_filtered():
    from app.facts import IndexedFactResult, project_indexed
    source={"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":"Customer commitment\nRespond to email complaints within 7 days"}]}
    segment={"id":"p1-1","kind":"pdf","target_ids":["p1:l1","p1:l2"],"records":[{"id":"p1:l1","unit_index":1,"text":"Customer commitment","cells":[]},{"id":"p1:l2","unit_index":1,"text":"Respond to email complaints within 7 days","cells":[]}]}
    fact={"label":"Customer commitment","source_label":"Customer commitment","label_ref":"p1:l1","label_refs":["p1:l1","p1:l1"],"value":"Respond to email complaints within 7 days","value_ref":"p1:l2","value_refs":["p1:l2","p1:l2"],"organization":None,"organization_ref":None,"unit":None,"unit_ref":None,"period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
    data={"document_type":"Policy","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    output=project_indexed(IndexedFactResult.model_validate(data),source,segment)
    assert not output["rejected_candidates"] and len(output["review_candidates"]) == 1
    candidate=output["review_candidates"][0]
    assert candidate["method"] == "local_exact_reference_recovery" and len(candidate["recovery_audit"]) == 2
    assert output["proposed_facts"][0]["value_refs"] == ["p1:l2","p1:l2"]


def test_local_revalidation_archives_before_after_once(tmp_path):
    config=replace(Settings(),data_dir=tmp_path,openai_api_key="test-key")
    with TestClient(create_app(config)) as client:
        doc=upload(client).json()["results"][0]["document"]
        store=Store(config); assert process_one(store); assert fact_queue.enqueue(store,doc["id"]) == "queued"
        job=fact_queue.claim(store)
        proposed={"document_type":"Workbook","organization":None,"organization_ref":None,"facts":[],"warnings":[],"coverage":"partial","coverage_note":"Review"}
        previous={"index_version":"source-index-v3","proposed_result":proposed,"result":{"facts":[],"review_candidates":[],"rejected_candidates":[],"review_issues":[],"warnings":[]},"validator_version":"earlier"}
        assert fact_queue.finish(store,job,output=previous)
        url=f'/api/documents/{doc["id"]}/facts/revalidate'
        assert client.post(url).status_code == 200
        history=client.get(f'/api/documents/{doc["id"]}/facts/history').json()["runs"]
        assert len(history)==1 and history[0]["job"]["audit"]["action"] == "local_revalidation"
        assert history[0]["job"]["audit"]["before_hash"] != history[0]["job"]["audit"]["after_hash"]
        assert json.loads(history[0]["job"]["segments"][0]["output_json"])["proposed_result"] == proposed
        assert client.post(url).status_code == 200
        assert len(client.get(f'/api/documents/{doc["id"]}/facts/history').json()["runs"]) == 1


@pytest.mark.parametrize("label,value", [("Website for service", "service.example.com"), ("Whatsapp Contact", "9594519122")])
def test_contact_identifiers_are_rejected_without_dropping_policy_statements(label, value):
    from app.facts import IndexedFactResult, project_indexed
    text=f"{label} {value}"
    source={"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":text}]}
    fact={"label":label,"source_label":label,"label_ref":"p1:l1","value":value,"value_ref":"p1:l1","organization":None,"organization_ref":None,"unit":None,"unit_ref":None,"period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
    data={"document_type":"Report","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    assert "outside" in project_indexed(IndexedFactResult.model_validate(data),source)["rejected_candidates"][0]["reason"]


def test_recovery_detects_single_line_and_split_line_duplicate_as_ambiguous():
    from app.facts import IndexedFactResult, project_indexed
    lines=["Current Month Bill", "Current Month", "Bill", "Value 42"]
    source={"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":"\n".join(lines)}]}
    segment={"id":"p1-1","kind":"pdf","target_ids":[f"p1:l{i}" for i in range(1,5)],"records":[{"id":f"p1:l{i}","unit_index":1,"text":line,"cells":[]} for i,line in enumerate(lines,1)]}
    fact={"label":"Current bill","source_label":"Current Month Bill","label_ref":"p1:l4","value":"42","value_ref":"p1:l4","organization":None,"organization_ref":None,"unit":None,"unit_ref":None,"period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
    data={"document_type":"Bill","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    result=project_indexed(IndexedFactResult.model_validate(data),source,segment)
    assert not result["review_candidates"] and result["rejected_candidates"]
    # One occurrence observed through many surrounding windows is still unique.
    lines=["Header","Current Month Bill","Other","Value 42"]
    source["units"][0]["text"]="\n".join(lines)
    segment["records"]=[{"id":f"p1:l{i}","unit_index":1,"text":line,"cells":[]} for i,line in enumerate(lines,1)]
    result=project_indexed(IndexedFactResult.model_validate(data),source,segment)
    assert len(result["review_candidates"]) == 1
    assert result["review_candidates"][0]["recovery_audit"][0]["recovered_ids"] == ["p1:l2"]


@pytest.mark.parametrize("field", ["source_label", "value"])
def test_whitespace_only_core_proposal_is_rejected_without_crash(field):
    from app.facts import IndexedFactResult, project_indexed
    source={"version":1,"kind":"pdf","summary":{"warnings":[]},"units":[{"index":1,"label":"Page 1","text":"Energy 42"}]}
    segment={"id":"p1-1","kind":"pdf","target_ids":["p1:l1"],"records":[{"id":"p1:l1","unit_index":1,"text":"Energy 42","cells":[]}]}
    fact={"label":"Energy","source_label":"Energy","label_ref":"p1:l99","value":"42","value_ref":"p1:l99","organization":None,"organization_ref":None,"unit":None,"unit_ref":None,"period":None,"period_ref":None,"currency":None,"currency_ref":None,"origin":"document_reported"}
    fact[field]="   "
    data={"document_type":"Report","organization":None,"organization_ref":None,"facts":[fact],"warnings":[],"coverage":"partial","coverage_note":"Review"}
    assert project_indexed(IndexedFactResult.model_validate(data),source,segment)["rejected_candidates"]
