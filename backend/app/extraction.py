"""Extract bounded, location-preserving source content; do not infer metrics."""
from datetime import date, datetime, time, timedelta
import math
from zipfile import ZipFile

from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from openpyxl.xml.functions import iterparse

import openpyxl
import pdfplumber

MAX_TEXT_CHARS = 10_000_000
MAX_CELLS = 100_000
MAX_ROWS = 100_000
MAX_COLUMNS = 256
MAX_GRID_POSITIONS = 2_000_000


def json_value(value):
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Workbook contains a non-finite numeric value; correct the source cell before retrying.")
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    raise ValueError("Workbook contains an unsupported cell value type.")


def check_sheet_dimensions(archive, sheet):
    """Reject inaccurate XLSX dimensions instead of silently dropping cells."""
    if sheet.max_row is None or sheet.max_column is None:
        raise ValueError("Workbook sheet has no declared dimensions; source extraction cannot safely scan it.")
    count = 0
    with archive.open(sheet._worksheet_path) as xml:
        for _, element in iterparse(xml):
            if element.tag.endswith("}c"):
                reference = element.get("r")
                try:
                    column, row = coordinate_from_string(reference)
                    column = column_index_from_string(column)
                except (TypeError, ValueError) as exc:
                    raise ValueError("Workbook contains a cell without a valid coordinate.") from exc
                if row > MAX_ROWS or column > MAX_COLUMNS:
                    raise ValueError("Workbook exceeds the 100,000 row or 256 column source limit.")
                if row < sheet.min_row or column < sheet.min_column or row > sheet.max_row or column > sheet.max_column:
                    raise ValueError("Workbook sheet dimensions omit cells; source extraction cannot safely continue.")
                count += 1
                if count > MAX_GRID_POSITIONS:
                    raise ValueError("Workbook exceeds the 2,000,000 cell-position scan limit.")
            element.clear()


LAYOUT_VERSION = "pdf-layout-v2"


def narrative_columns(page):
    """Detect a persistent central gutter on narrative pages; preserve header/footer separately."""
    words = page.extract_words()
    body_top, body_bottom = page.height * 0.10, page.height * 0.70
    if any(word["top"] < boundary < word["bottom"] for word in words for boundary in (body_top, body_bottom)):
        return None
    body = [word for word in words if body_top <= word["top"] and word["bottom"] <= body_bottom]
    if len(body) < 150 or sum(any(char.isdigit() for char in word["text"]) for word in body) / len(body) >= 0.05:
        return None
    candidates = []
    for step in range(35, 66):
        split = page.width * step / 100
        crossing = [word for word in body if word["x0"] < split < word["x1"]]
        left = [word for word in body if word["x1"] <= split - 2]
        right = [word for word in body if word["x0"] >= split + 2]
        if not crossing and len(left) >= 50 and len(right) >= 50:
            # Prefer the widest margin between the left and right word edges.
            gap = min(word["x0"] for word in right) - max(word["x1"] for word in left)
            if gap >= 4:
                candidates.append((gap, split))
    if not candidates:
        return None
    _, split = max(candidates)
    boundaries = [("Header", (0, 0, page.width, body_top)),
                  ("Left column", (0, body_top, split, body_bottom)),
                  ("Right column", (split, body_top, page.width, body_bottom)),
                  ("Footer", (0, body_bottom, page.width, page.height))]
    blocks = []
    for label, bbox in boundaries:
        text = page.crop(bbox).extract_text() or ""
        if text.strip():
            blocks.append({"label": label, "bbox": list(bbox), "text": text})
    return blocks if len(blocks) >= 2 else None


def extract_source(path, kind):
    warnings = []
    units = []
    if kind == "pdf":
        total = 0
        column_pages = 0
        with pdfplumber.open(path) as pdf:
            for number, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                total += len(text)
                if total > MAX_TEXT_CHARS:
                    raise ValueError("PDF source text exceeds the 10,000,000 character limit.")
                if not text.strip():
                    warnings.append(f"Page {number} has no extractable text; OCR may be needed.")
                try:
                    blocks = narrative_columns(page)
                except (ValueError, TypeError, KeyError, ArithmeticError):
                    blocks = None
                    warnings.append(f"Page {number}: reading layout could not be determined; original page text is preserved. Review reading order.")
                if blocks:
                    column_pages += 1
                    total += sum(len(block["text"]) for block in blocks)
                    if total > MAX_TEXT_CHARS:
                        raise ValueError("PDF source text exceeds the 10,000,000 character limit.")
                units.append({"index": number, "label": f"Page {number}", "text": text, "reading_blocks": blocks})
        summary = {"unit_count": len(units), "text_characters": total, "warnings": warnings, "layout_version": LAYOUT_VERSION, "column_pages": column_pages}
    elif kind == "xlsx":
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        cached = None
        try:
            cached = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
            count = 0
            positions = 0
            text_size = 0
            for index, sheet in enumerate(formulas, 1):
                with ZipFile(path) as archive:
                    check_sheet_dimensions(archive, sheet)
                if sheet.max_row is None or sheet.max_column is None:
                    raise ValueError("Workbook sheet has no declared dimensions; source extraction cannot safely scan it.")
                if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS:
                    raise ValueError("Workbook exceeds the 100,000 row or 256 column source limit.")
                positions += sheet.max_row * sheet.max_column
                if positions > MAX_GRID_POSITIONS:
                    raise ValueError("Workbook exceeds the 2,000,000 cell-position scan limit.")
                cells = []
                missing_caches = 0
                cached_sheet = cached[sheet.title]
                for row, cached_row in zip(sheet.iter_rows(), cached_sheet.iter_rows(), strict=True):
                    for cell, cached_cell in zip(row, cached_row, strict=True):
                        if cell.value is None:
                            continue
                        count += 1
                        if count > MAX_CELLS:
                            raise ValueError("Workbook exceeds the 100,000 nonempty cell source limit.")
                        is_formula = cell.data_type == "f"
                        if is_formula and not isinstance(cell.value, str):
                            raise ValueError("Workbook contains an unsupported formula representation.")
                        value = json_value(cached_cell.value if is_formula else cell.value)
                        text_size += len(str(cell.value)) + (len(str(value)) if is_formula and value is not None else 0)
                        if text_size > MAX_TEXT_CHARS:
                            raise ValueError("Workbook source content exceeds the 10,000,000 character limit.")
                        if is_formula and cached_cell.value is None:
                            missing_caches += 1
                        if cell.data_type == "e" or (is_formula and cached_cell.data_type == "e"):
                            warnings.append(f"Sheet {sheet.title}, cell {cell.coordinate}: spreadsheet error value preserved; review the source cell.")
                        cells.append({"coordinate": cell.coordinate, "type": cell.data_type,
                                      "value": value, "number_format": cell.number_format, "formula": cell.value if is_formula else None,
                                      "cached_type": cached_cell.data_type if is_formula else None,
                                      "cached_value": json_value(cached_cell.value) if is_formula else None})
                if missing_caches:
                    warnings.append(f"Sheet {sheet.title}: {missing_caches} formula cache(s) missing; formulas were not recalculated.")
                units.append({"index": index, "label": sheet.title, "cells": cells})
            summary = {"unit_count": len(units), "cell_count": count, "warnings": warnings}
        finally:
            formulas.close()
            if cached is not None:
                cached.close()
    else:
        raise ValueError("Unsupported file type.")
    return {"version": 1, "kind": kind, "summary": summary, "units": units}
