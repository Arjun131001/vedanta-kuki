"""File readability checks only; no sustainability metric extraction."""
from pathlib import Path
from zipfile import ZipFile, BadZipFile

import openpyxl
import pdfplumber


def validate_file(path: Path, kind: str):
    if kind == "pdf":
        with path.open("rb") as stream:
            if stream.read(5) != b"%PDF-":
                raise ValueError("The file does not have a PDF signature.")
        with pdfplumber.open(path) as pdf:
            if not pdf.pages:
                raise ValueError("The PDF has no pages.")
            if len(pdf.pages) > 2000:
                raise ValueError("PDF exceeds the 2,000-page validation limit.")
            # Walk page dimensions to validate the page tree without extracting content.
            for page in pdf.pages:
                if page.width <= 0 or page.height <= 0:
                    raise ValueError("The PDF contains invalid page dimensions.")
            return {"page_count": len(pdf.pages), "check": "PDF readability", "metrics_extracted": False}
    if kind == "xlsx":
        try:
            with ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(item.file_size for item in entries) > 100 * 1024 * 1024:
                    raise ValueError("Workbook expands beyond the 100 MB validation limit.")
                if any(item.flag_bits & 1 for item in entries):
                    raise ValueError("Encrypted workbooks are not supported.")
                if archive.testzip() is not None:
                    raise ValueError("The workbook archive is damaged.")
        except BadZipFile as exc:
            raise ValueError("The file is not a valid XLSX workbook.") from exc
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if not workbook.sheetnames:
                raise ValueError("The workbook has no worksheets.")
            # Force a bounded read from every sheet, in addition to workbook metadata.
            for sheet in workbook:
                next(sheet.iter_rows(max_row=1, max_col=20, values_only=True), None)
            return {"sheet_count": len(workbook.sheetnames), "check": "Workbook readability", "metrics_extracted": False}
        finally:
            workbook.close()
    raise ValueError("Unsupported file type.")
